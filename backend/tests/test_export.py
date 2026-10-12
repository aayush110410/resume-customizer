import io

import docx
import fitz
import pytest
from fastapi.testclient import TestClient

from app.export.docx import render_letter_docx, render_resume_docx
from app.export.pdf import render_letter_pdf, render_resume_pdf
from app.export.txt import render_txt
from app.guard.grounding import verify_quote
from app.knowledge.loader import load_knowledge
from app.parsing.extract import parse_resume
from app.resume_ops import assign_ids, iter_bullets


@pytest.fixture
def doc(sample_doc):
    return assign_ids(sample_doc)


@pytest.fixture
def kb():
    return load_knowledge()


def test_pdf_round_trip_is_ats_clean(doc, kb):
    r = parse_resume(render_resume_pdf(doc, kb.regions["india"]), 5)
    assert not r.layout.multi_column and r.layout.hidden_chars == 0 and not r.layout.has_tables
    assert "Arjun Mehta" in r.text and "EXPERIENCE" in r.text.upper()
    for _, _, b in iter_bullets(doc):
        assert verify_quote(b.text, r.text), b.text


def test_paper_size_by_region(doc, kb):
    a4 = fitz.open(stream=render_resume_pdf(doc, kb.regions["india"]), filetype="pdf")[0].rect.width
    letter = fitz.open(stream=render_resume_pdf(doc, kb.regions["international"]), filetype="pdf")[0].rect.width
    assert round(a4) == 595 and round(letter) == 612


def test_markup_characters_render_literally(doc, kb):
    doc.summary = "Scaled #payments by $5M with *care* and _focus_ @team <fast> [ok] = \\done"
    text = parse_resume(render_resume_pdf(doc, kb.regions["india"]), 5).text
    assert verify_quote(doc.summary, text)


def test_docx_contains_headings_and_bullets(doc, kb):
    d = docx.Document(io.BytesIO(render_resume_docx(doc, kb.regions["india"])))
    text = "\n".join(p.text for p in d.paragraphs)
    assert "EXPERIENCE" in text and doc.experience[0].bullets[0].text in text and "arjun@example.com" in text


def test_txt_has_contact_then_sections(doc):
    txt = render_txt(doc)
    assert txt.startswith("Arjun Mehta\narjun@example.com") and txt.index("SUMMARY") < txt.index("EXPERIENCE")


def test_letters_render(doc, kb):
    letter = "Dear Hiring Manager,\n\nI build payment systems.\n\nRegards,\nArjun Mehta"
    pdf_text = "".join(
        p.get_text()
        for p in fitz.open(stream=render_letter_pdf(letter, doc.contact, kb.regions["india"]), filetype="pdf")
    )
    assert "I build payment systems." in pdf_text
    d = docx.Document(io.BytesIO(render_letter_docx(letter, doc.contact)))
    assert any(p.text == "I build payment systems." for p in d.paragraphs)


def test_export_endpoint(doc):
    from app.main import create_app

    client = TestClient(create_app())
    r = client.post(
        "/api/export", json={"kind": "resume", "format": "pdf", "region": "india", "resume": doc.model_dump()}
    )
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"] == 'attachment; filename="Arjun_Mehta_Resume.pdf"'
    cl = client.post(
        "/api/export",
        json={
            "kind": "cover_letter",
            "format": "txt",
            "region": "india",
            "cover_letter": "Hi",
            "contact": doc.contact.model_dump(),
        },
    )
    assert cl.text == "Hi" and cl.headers["content-disposition"].endswith('Arjun_Mehta_Cover_Letter.txt"')
    bad = client.post("/api/export", json={"kind": "resume", "format": "pdf", "region": "india"})
    assert bad.status_code == 422 and bad.json()["error"]["code"] == "INVALID_EXPORT"

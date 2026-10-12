import pytest
from fastapi.testclient import TestClient

from app.schemas import Bullet, Contact, Entry, ResumeDoc, SkillGroup


@pytest.fixture
def sample_doc() -> ResumeDoc:
    return ResumeDoc(
        contact=Contact(name="Arjun Mehta", email="arjun@example.com", phone="+91 98765 43210", location="Bengaluru"),
        summary="Backend engineer with 4 years building payment APIs.",
        skills=[SkillGroup(name="Languages", items=["Python", "Go"])],
        experience=[
            Entry(
                title="Software Engineer",
                organization="Razorpay",
                start="Jan 2022",
                end="Present",
                bullets=[
                    Bullet(text="Reduced p99 latency by 30% for 2M daily payments"),
                    Bullet(text="Built REST APIs"),
                ],
            ),
            Entry(
                title="Associate Engineer",
                organization="Infosys",
                start="Jul 2020",
                end="Dec 2021",
                bullets=[Bullet(text="Maintained billing jobs"), Bullet(text="Wrote SQL reports")],
            ),
        ],
        projects=[Entry(title="Ledger CLI", organization="Personal", bullets=[Bullet(text="Open-source ledger tool")])],
        education=[Entry(title="B.Tech, Computer Science", organization="NIT Trichy", end="2020")],
        achievements=[Bullet(text="Winner, Smart India Hackathon 2019"), Bullet(text="AWS Certified Developer")],
    )


@pytest.fixture
def client() -> TestClient:
    from app.main import create_app

    return TestClient(create_app())


# --- file fixtures (generated, no binaries committed) -------------------------------------------------------
import io  # noqa: E402

import docx  # noqa: E402
import fitz  # noqa: E402

LOREM = (
    "Software Engineer at Razorpay, Bengaluru. Reduced p99 latency by 30% for 2M daily payments. "
    "Built REST APIs in Python and Go. Led migration to Kubernetes. Mentored four engineers. "
)


def _pdf(build) -> bytes:
    doc = fitz.open()
    build(doc)
    return doc.tobytes()


@pytest.fixture
def simple_pdf() -> bytes:
    def build(doc):
        page = doc.new_page()
        y = 60
        for heading in ("Summary", "Experience", "Education", "Skills"):
            page.insert_text((50, y), heading, fontsize=14)
            page.insert_textbox(fitz.Rect(50, y + 8, 550, y + 90), LOREM, fontsize=10)
            y += 110

    return _pdf(build)


@pytest.fixture
def two_column_pdf() -> bytes:
    def build(doc):
        page = doc.new_page()
        for i in range(4):
            y = 60 + i * 120
            page.insert_textbox(fitz.Rect(40, y, 270, y + 100), "Left " + LOREM, fontsize=9)
            page.insert_textbox(fitz.Rect(320, y, 560, y + 100), "Right " + LOREM, fontsize=9)

    return _pdf(build)


@pytest.fixture
def hidden_text_pdf() -> bytes:
    def build(doc):
        page = doc.new_page()
        page.insert_textbox(fitz.Rect(50, 50, 550, 300), LOREM * 2, fontsize=10)
        page.insert_text((50, 700), "Kubernetes Terraform", fontsize=10, color=(1, 1, 1))

    return _pdf(build)


@pytest.fixture
def scanned_pdf() -> bytes:
    def build(doc):
        page = doc.new_page()
        pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 60, 60), 0)
        pix.clear_with(200)
        page.insert_image(fitz.Rect(50, 50, 500, 700), pixmap=pix)

    return _pdf(build)


@pytest.fixture
def encrypted_pdf() -> bytes:
    doc = fitz.open()
    doc.new_page().insert_textbox(fitz.Rect(50, 50, 550, 400), LOREM * 3)
    buf = io.BytesIO()
    doc.save(buf, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="x", owner_pw="y")
    return buf.getvalue()


@pytest.fixture
def six_page_pdf() -> bytes:
    def build(doc):
        for _ in range(6):
            doc.new_page().insert_textbox(fitz.Rect(50, 50, 550, 400), LOREM * 2)

    return _pdf(build)


def _docx(build) -> bytes:
    d = docx.Document()
    build(d)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


@pytest.fixture
def simple_docx() -> bytes:
    def build(d):
        for heading in ("Summary", "Experience", "Skills"):
            d.add_heading(heading, level=2)
            d.add_paragraph(LOREM)

    return _docx(build)


@pytest.fixture
def docx_with_table() -> bytes:
    def build(d):
        d.add_heading("Experience", level=2)
        d.add_paragraph(LOREM * 2)
        t = d.add_table(rows=2, cols=3)
        for i, v in enumerate(("Degree", "Institute", "CGPA", "B.Tech", "NIT Trichy", "8.4/10")):
            t.cell(i // 3, i % 3).text = v
        d.sections[0].header.paragraphs[0].text = "Arjun Mehta | +91 98765 43210"

    return _docx(build)

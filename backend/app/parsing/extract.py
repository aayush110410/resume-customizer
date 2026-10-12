"""PDF/DOCX → text plus the layout signals ATS parsers trip over (columns, tables, images, hidden text)."""

import io
import zipfile
from typing import Literal

import docx
import pymupdf as fitz

from app.schemas import LayoutReport, ParsedResume
from app.text_utils import estimate_pages

MIN_TEXT_CHARS = 200
HIDDEN_COLOR = 0xFFFFFF
TINY_FONT_PT = 4


class ParseError(Exception):
    def __init__(self, code: str):
        super().__init__(code)
        self.code = code


def detect_type(data: bytes) -> Literal["pdf", "docx"]:
    if not data:
        raise ParseError("EMPTY_FILE")
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" in z.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
    raise ParseError("INVALID_FILE_TYPE")


def parse_resume(data: bytes, max_pages: int) -> ParsedResume:
    kind = detect_type(data)
    text, layout = _parse_pdf(data, max_pages) if kind == "pdf" else _parse_docx(data, max_pages)
    if len(text.strip()) < MIN_TEXT_CHARS:
        raise ParseError("SCANNED_PDF" if kind == "pdf" else "EMPTY_FILE")
    return ParsedResume(text=text, layout=layout)


def _is_multi_column(blocks: list[tuple[float, float, float, float]], width: float, height: float) -> bool:
    left = [b for b in blocks if b[2] < 0.55 * width]
    right = [b for b in blocks if b[0] > 0.45 * width]
    if len(left) < 3 or len(right) < 3:
        return False
    overlap = min(max(b[3] for b in left), max(b[3] for b in right)) - max(
        min(b[1] for b in left), min(b[1] for b in right)
    )
    return overlap > 0.3 * height


def _parse_pdf(data: bytes, max_pages: int) -> tuple[str, LayoutReport]:
    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception as e:  # PyMuPDF raises several types for malformed files
        raise ParseError("CORRUPT_FILE") from e
    if doc.needs_pass:
        raise ParseError("ENCRYPTED_PDF")
    if doc.page_count == 0:
        raise ParseError("CORRUPT_FILE")
    if doc.page_count > max_pages:
        raise ParseError("TOO_MANY_PAGES")

    blocks_text: list[str] = []
    hidden = images = 0
    multi_column = has_tables = False
    for page in doc:
        rects: list[tuple[float, float, float, float]] = []
        for block in page.get_text("dict")["blocks"]:
            if block.get("type") != 0:
                continue
            lines: list[str] = []
            for line in block["lines"]:
                visible = []
                for span in line["spans"]:
                    if span["color"] == HIDDEN_COLOR or span["size"] < TINY_FONT_PT:
                        hidden += len(span["text"].strip())
                    else:
                        visible.append(span["text"])
                if "".join(visible).strip():
                    lines.append("".join(visible))
            if lines:
                blocks_text.append("\n".join(lines))
                rects.append(tuple(block["bbox"]))
        multi_column = multi_column or _is_multi_column(rects, page.rect.width, page.rect.height)
        has_tables = has_tables or bool(page.find_tables().tables)
        images += len(page.get_images())

    text = "\n".join(blocks_text)
    return text, LayoutReport(
        file_type="pdf",
        pages=doc.page_count,
        char_count=len(text),
        word_count=len(text.split()),
        multi_column=multi_column,
        has_tables=has_tables,
        image_count=images,
        hidden_chars=hidden,
    )


def _is_hidden_run(run) -> bool:
    font = run.font
    color = font.color.rgb if font.color is not None and font.color.type is not None else None
    return bool(font.hidden) or str(color) == "FFFFFF" or (font.size is not None and font.size.pt < TINY_FONT_PT)


def _parse_docx(data: bytes, max_pages: int) -> tuple[str, LayoutReport]:
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError("CORRUPT_FILE") from e

    hidden = 0
    parts: list[str] = []
    for p in d.paragraphs:
        visible = []
        for run in p.runs:
            if _is_hidden_run(run):
                hidden += len(run.text.strip())
            else:
                visible.append(run.text)
        if "".join(visible).strip():
            parts.append("".join(visible))
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(dict.fromkeys(cells)))
    header_footer = any(p.text.strip() for s in d.sections for part in (s.header, s.footer) for p in part.paragraphs)
    text = "\n".join(parts)
    words = len(text.split())
    pages = estimate_pages(words)
    if pages > max_pages:
        raise ParseError("TOO_MANY_PAGES")
    return text, LayoutReport(
        file_type="docx",
        pages=pages,
        char_count=len(text),
        word_count=words,
        has_tables=bool(d.tables),
        image_count=len(d.inline_shapes),
        hidden_chars=hidden,
        header_footer_text=header_footer,
    )

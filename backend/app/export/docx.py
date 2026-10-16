"""DOCX rendering with python-docx: plain paragraphs and the built-in List Bullet style (ATS-friendly)."""

import io
import re

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.export.pdf import paragraphs
from app.knowledge.models import RegionPack
from app.schemas import Bullet, Contact, Entry, ResumeDoc

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _clean(value):
    """XML 1.0 forbids most control characters; python-docx raises on them. Cleans strings inside dicts/lists too."""
    if isinstance(value, str):
        return _CONTROL.sub("", value)
    if isinstance(value, dict):
        return {k: _clean(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_clean(v) for v in value]
    return value


def _contact_line(c: Contact) -> str:
    return "  |  ".join(d for d in (c.email, c.phone, c.location, *c.links) if d)


def _header(d: docx.Document, c: Contact, align=WD_ALIGN_PARAGRAPH.CENTER) -> None:
    name = d.add_paragraph()
    name.alignment = align
    run = name.add_run(c.name)
    run.bold, run.font.size = True, Pt(18)
    d.add_paragraph(_contact_line(c)).alignment = align


def _section(d: docx.Document, title: str) -> None:
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(title.upper())
    run.bold, run.font.size = True, Pt(11)


def _bullets(d: docx.Document, bullets: list[Bullet]) -> None:
    for b in bullets:
        d.add_paragraph(b.text, style="List Bullet")


def _entry(d: docx.Document, e: Entry) -> None:
    p = d.add_paragraph()
    p.add_run(e.title).bold = True
    dates = " – ".join(x for x in (e.start, e.end) if x)
    rest = " | ".join(x for x in (e.location, dates) if x)
    p.add_run((f", {e.organization}" if e.organization else "") + (f" | {rest}" if rest else ""))
    _bullets(d, e.bullets)


def _save(d: docx.Document) -> bytes:
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def render_resume_docx(doc: ResumeDoc, region: RegionPack) -> bytes:
    doc = ResumeDoc.model_validate(_clean(doc.model_dump()))
    d = docx.Document()
    d.styles["Normal"].font.size = Pt(10.5)
    _header(d, doc.contact)
    if doc.summary:
        _section(d, "Summary")
        d.add_paragraph(doc.summary)
    if doc.skills:
        _section(d, "Skills")
        for g in doc.skills:
            p = d.add_paragraph()
            p.add_run(f"{g.name}: ").bold = True
            p.add_run(", ".join(g.items))
    for key in ("experience", "projects", "education"):
        entries = getattr(doc, key)
        if entries:
            _section(d, key)
            for e in entries:
                _entry(d, e)
    if doc.achievements:
        _section(d, "Achievements")
        _bullets(d, doc.achievements)
    for o in doc.other_sections:
        _section(d, o.name)
        _bullets(d, o.bullets)
    return _save(d)


def render_letter_docx(text: str, contact: Contact) -> bytes:
    text, contact = _clean(text), Contact.model_validate(_clean(contact.model_dump()))
    d = docx.Document()
    _header(d, contact, align=WD_ALIGN_PARAGRAPH.LEFT)
    for block in paragraphs(text):
        p = d.add_paragraph()
        for i, line in enumerate(block.split("\n")):
            if i:
                p.runs[-1].add_break()
            p.add_run(line).bold = block.startswith("Subject:")
    return _save(d)
